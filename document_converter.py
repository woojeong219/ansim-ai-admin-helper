import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path


OPERATIONS = [
    "엑셀 → PDF",
    "Word → PDF",
    "이미지 → PDF",
    "HWP → PDF",
    "HWP → Word",
    "PDF → Excel",
    "PDF → Word",
    "PDF → 이미지",
    "PDF → HWP (지원 안내)",
    "PDF 여러 개 합치기",
]


ACCEPTED_TYPES = {
    "엑셀 → PDF": ["xlsx", "xls"],
    "Word → PDF": ["docx", "doc"],
    "이미지 → PDF": ["png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff"],
    "HWP → PDF": ["hwp", "hwpx"],
    "HWP → Word": ["hwp", "hwpx"],
    "PDF → Excel": ["pdf"],
    "PDF → Word": ["pdf"],
    "PDF → 이미지": ["pdf"],
    "PDF → HWP (지원 안내)": ["pdf"],
    "PDF 여러 개 합치기": ["pdf"],
}


def _safe_name(name):
    return Path(name).name.replace("\x00", "")


def _package_outputs(outputs):
    if len(outputs) == 1:
        return outputs[0]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data, _mime in outputs:
            archive.writestr(name, data)
    return "변환결과.zip", buffer.getvalue(), "application/zip"


def _libreoffice_convert(upload, target_extension):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("문서 변환 엔진을 찾지 못했습니다. GitHub에 packages.txt가 올라갔는지 확인해 주세요.")
    with tempfile.TemporaryDirectory(prefix="document-convert-") as temp_dir:
        work = Path(temp_dir)
        input_path = work / _safe_name(upload.name)
        input_path.write_bytes(upload.getvalue())
        profile = work / "profile"
        command = [
            executable,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--nofirststartwizard",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            target_extension,
            "--outdir",
            str(work),
            str(input_path),
        ]
        completed = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
        output_path = work / f"{input_path.stem}.{target_extension}"
        if completed.returncode != 0 or not output_path.exists():
            detail = (completed.stderr or completed.stdout or "지원하지 않는 파일 형식일 수 있습니다.").strip()
            raise RuntimeError(f"변환하지 못했습니다: {detail[:300]}")
        mime = "application/pdf" if target_extension == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        return output_path.name, output_path.read_bytes(), mime


def _images_to_pdf(uploads):
    from PIL import Image

    images = []
    for upload in uploads:
        image = Image.open(io.BytesIO(upload.getvalue()))
        if image.mode in ("RGBA", "LA"):
            background = Image.new("RGB", image.size, "white")
            alpha = image.getchannel("A")
            background.paste(image, mask=alpha)
            image = background
        else:
            image = image.convert("RGB")
        images.append(image)
    if not images:
        raise ValueError("이미지 파일을 하나 이상 올려주세요.")
    buffer = io.BytesIO()
    images[0].save(buffer, format="PDF", save_all=True, append_images=images[1:])
    return "이미지_변환결과.pdf", buffer.getvalue(), "application/pdf"


def _merge_pdfs(uploads):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for upload in uploads:
        reader = PdfReader(io.BytesIO(upload.getvalue()))
        for page in reader.pages:
            writer.add_page(page)
    buffer = io.BytesIO()
    writer.write(buffer)
    return "PDF_합치기_결과.pdf", buffer.getvalue(), "application/pdf"


def _pdf_to_images(upload):
    import fitz

    document = fitz.open(stream=upload.getvalue(), filetype="pdf")
    outputs = []
    for page_number, page in enumerate(document, start=1):
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        outputs.append((f"{Path(upload.name).stem}_{page_number:03d}.png", pixmap.tobytes("png"), "image/png"))
    document.close()
    return _package_outputs(outputs)


def _pdf_bytes_to_word_layout(pdf_bytes, output_stem):
    import fitz
    from docx import Document
    from docx.shared import Inches, Pt

    pdf = fitz.open(stream=pdf_bytes, filetype="pdf")
    if len(pdf) == 0:
        raise ValueError("PDF에 변환할 페이지가 없습니다.")
    document = Document()
    first_rect = pdf[0].rect
    section = document.sections[0]
    section.page_width = Inches(first_rect.width / 72)
    section.page_height = Inches(first_rect.height / 72)
    for margin_name in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, margin_name, Inches(0.1))

    for page_number, page in enumerate(pdf):
        if page_number:
            document.add_page_break()
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image_stream = io.BytesIO(pixmap.tobytes("png"))
        page_width = page.rect.width / 72
        page_height = page.rect.height / 72
        scale = min((section.page_width.inches - 0.2) / page_width, (section.page_height.inches - 0.2) / page_height)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        run = paragraph.add_run()
        run.add_picture(image_stream, width=Inches(page_width * scale), height=Inches(page_height * scale))
    pdf.close()
    buffer = io.BytesIO()
    document.save(buffer)
    return f"{output_stem}_원본모습.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _pdf_to_word_layout(upload):
    return _pdf_bytes_to_word_layout(upload.getvalue(), Path(upload.name).stem)


def _pdf_to_excel_layout(upload):
    import fitz
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ExcelImage

    pdf = fitz.open(stream=upload.getvalue(), filetype="pdf")
    if len(pdf) == 0:
        raise ValueError("PDF에 변환할 페이지가 없습니다.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    image_streams = []
    for page_number, page in enumerate(pdf, start=1):
        worksheet = workbook.create_sheet(f"{page_number}페이지")
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        image_stream = io.BytesIO(pixmap.tobytes("png"))
        image_streams.append(image_stream)
        image = ExcelImage(image_stream)
        worksheet.add_image(image, "A1")
        worksheet.sheet_view.showGridLines = False
    pdf.close()
    buffer = io.BytesIO()
    workbook.save(buffer)
    return f"{Path(upload.name).stem}_원본모습.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _pdf_to_word(upload):
    import fitz
    from docx import Document

    pdf = fitz.open(stream=upload.getvalue(), filetype="pdf")
    document = Document()
    document.add_heading(Path(upload.name).stem, level=1)
    for page_number, page in enumerate(pdf, start=1):
        if page_number > 1:
            document.add_page_break()
        document.add_heading(f"{page_number}페이지", level=2)
        blocks = page.get_text("blocks", sort=True)
        page_text = False
        for block in blocks:
            text = block[4].strip()
            if text:
                document.add_paragraph(text)
                page_text = True
        if not page_text:
            document.add_paragraph("[추출할 수 있는 텍스트가 없습니다. 스캔 문서는 OCR이 필요합니다.]")
    pdf.close()
    buffer = io.BytesIO()
    document.save(buffer)
    return f"{Path(upload.name).stem}.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _pdf_to_excel(upload):
    import pandas as pd
    import pdfplumber

    tables = []
    text_rows = []
    with pdfplumber.open(io.BytesIO(upload.getvalue())) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            for table_number, table in enumerate(page.extract_tables(), start=1):
                if table and len(table) > 1:
                    headers = [str(value).strip() if value not in (None, "") else f"열{index + 1}" for index, value in enumerate(table[0])]
                    tables.append((f"P{page_number}_표{table_number}", pd.DataFrame(table[1:], columns=headers)))
            text = page.extract_text() or ""
            for line in text.splitlines():
                text_rows.append({"페이지": page_number, "추출텍스트": line})
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        if tables:
            for sheet_name, frame in tables[:50]:
                frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)
        else:
            pd.DataFrame(text_rows or [{"페이지": "", "추출텍스트": "추출할 표나 텍스트가 없습니다."}]).to_excel(
                writer, sheet_name="추출내용", index=False
            )
    return f"{Path(upload.name).stem}.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def convert_uploads(operation, uploads, conversion_mode="editable"):
    if not uploads:
        raise ValueError("변환할 파일을 올려주세요.")
    if operation == "이미지 → PDF":
        return _images_to_pdf(uploads)
    if operation == "PDF 여러 개 합치기":
        if len(uploads) < 2:
            raise ValueError("합칠 PDF 파일을 두 개 이상 올려주세요.")
        return _merge_pdfs(uploads)

    outputs = []
    for upload in uploads:
        if operation in ("엑셀 → PDF", "Word → PDF", "HWP → PDF"):
            outputs.append(_libreoffice_convert(upload, "pdf"))
        elif operation == "HWP → Word":
            if conversion_mode == "layout":
                _pdf_name, pdf_bytes, _pdf_mime = _libreoffice_convert(upload, "pdf")
                outputs.append(_pdf_bytes_to_word_layout(pdf_bytes, Path(upload.name).stem))
            else:
                outputs.append(_libreoffice_convert(upload, "docx"))
        elif operation == "PDF → Word":
            outputs.append(_pdf_to_word_layout(upload) if conversion_mode == "layout" else _pdf_to_word(upload))
        elif operation == "PDF → Excel":
            outputs.append(_pdf_to_excel_layout(upload) if conversion_mode == "layout" else _pdf_to_excel(upload))
        elif operation == "PDF → 이미지":
            outputs.append(_pdf_to_images(upload))
    return _package_outputs(outputs)
